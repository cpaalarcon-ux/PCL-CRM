"""Connected staging suite. Only fictional QA accounts supplied by a private fixture.
Never put credentials, production records or workbooks in this file.
"""
import json,sys,pathlib,urllib.request,urllib.error,uuid,datetime,io,base64,zipfile
from openpyxl import Workbook
from reportlab.pdfgen import canvas
fixture=pathlib.Path(sys.argv[1]); ctx=json.loads(fixture.read_text()); state=fixture.with_suffix('.suite.json')
URL='https://yfjiculexbzpkvtapudk.supabase.co/functions/v1/staff-workspace-v7'
origin='https://backend-staging-foundation.pcl-crm.pages.dev'
def call(body,token=None,expect=200):
    req=urllib.request.Request(URL,data=json.dumps(body).encode(),headers={'Content-Type':'application/json','Origin':origin,**({'Authorization':'Bearer '+token} if token else {})})
    try:
        with urllib.request.urlopen(req,timeout=90) as r: status=r.status;out=json.load(r)
    except urllib.error.HTTPError as e:status=e.code;out=json.load(e)
    assert status==expect,(body['action'],status,out.get('error','Unexpected response'))
    print('PASS',body['action'],status,flush=True);return out

def multipart(fields,data,name,mime,token,kind=None,expect=200):
    boundary='S7'+uuid.uuid4().hex;parts=[]
    for k,v in fields.items():parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
    parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\nContent-Type: {mime}\r\n\r\n'.encode()+data+b'\r\n');parts.append(f'--{boundary}--\r\n'.encode())
    req=urllib.request.Request(URL,data=b''.join(parts),headers={'Origin':origin,'Authorization':'Bearer '+token,'Content-Type':'multipart/form-data; boundary='+boundary,**({'x-erp-upload-kind':kind} if kind else {})})
    try:
        with urllib.request.urlopen(req,timeout=90) as r:status=r.status;out=json.load(r)
    except urllib.error.HTTPError as e:status=e.code;out=json.load(e)
    assert status==expect,(name,status,out.get('error','Unexpected response'))
    print('PASS',fields.get('action','private upload'),status,flush=True);return out

def book(rows):
    b=Workbook();s=b.active;s.title='Data'
    for row in rows:s.append(row)
    stream=io.BytesIO();b.save(stream);return stream.getvalue()
xlsx='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
P=call({'action':'login','login_id':ctx['login_id'],'password':ctx['password']})['access_token']
A=call({'action':'login','login_id':ctx['ASSISTANT']['login_id'],'password':ctx['ASSISTANT']['password']})['access_token']
S=call({'action':'login','login_id':ctx['SUPERVISOR']['login_id'],'password':ctx['SUPERVISOR']['password']})['access_token']
D=json.loads(state.read_text()) if state.exists() else {}
def save():state.write_text(json.dumps(D));state.chmod(0o600)
def w(token=P):return call({'action':'workspace'},token)
def get(key,idkey,value,token=P):return next(x for x in w(token)['suite'][key] if x[idkey]==value)
firm=ctx['firm'];today=str(datetime.date.today());suffix=D.get('suffix') or uuid.uuid4().hex[:7];D['suffix']=suffix;save()
if 'client' not in D:
    D['client']=call({'action':'client_save','company_name':'S7 QA fictional '+suffix,'kra_pin':'P'+str(int(suffix,16)).zfill(9)+'Q','primary_sector':'SME/Startup','branch_location':'Nairobi','client_category':'Corporate','firm_ids':[firm]},P)['client_id'];save()
cid=D['client']
if 'project' not in D:
    D['project']=call({'action':'project_save','client_id':cid,'firm_id':firm,'service_code':'TAX-COMP','frequency':'Weekly + Monthly','pack':'Tax Compilation','starts_on':'2026-09-01','ends_on':'2026-12-31','assigned_to':ctx['ASSISTANT']['staff_id'],'reviewer_id':ctx['SUPERVISOR']['staff_id'],'obligations':['VAT','PAYE']},P)['project_id'];save()
pid=D['project']
if 'activated' not in D:
    call({'action':'engagement_activate','project_id':pid},A,403)
    call({'action':'engagement_activate','project_id':pid},P,400)
    out=io.BytesIO();pdf=canvas.Canvas(out);pdf.drawString(72,720,'Fictional QA evidence - no real client information');pdf.save();data=out.getvalue()
    permanent=multipart({'client_id':cid,'firm_id':firm,'category':'Permanent File'},data,'QA_identity.pdf','application/pdf',P)['document_id']
    letter=multipart({'client_id':cid,'firm_id':firm,'project_id':pid,'category':'Signed Engagement Letters'},data,'QA_signed_letter.pdf','application/pdf',P)['document_id']
    call({'action':'compliance_save','client_id':cid,'entity_type':'Sole proprietor','risk_rating':'Low','is_pep':False,'conflict_checked':True,'sanctions_date':today,'sanctions_result':'Fictional QA screening only','review_due':'2027-09-30','incorporation_document':permanent,'pin_document':permanent,'notes':'QA only'},P)
    k=get('compliance','client_id',cid);call({'action':'compliance_verify','client_id':cid,'revision':k['revision']},P)
    call({'action':'engagement_save','project_id':pid,'signed_letter_id':letter,'letter_expires_on':'2027-09-30','stage':'Accepted'},P)
    before={t['task_id'] for t in w()['tasks'] if t['project_id']==pid}
    call({'action':'engagement_activate','project_id':pid},P)
    ws=w();assert next(p for p in ws['projects'] if p['project_id']==pid)['status']=='Active';assert before.issubset({t['task_id'] for t in ws['tasks']})
    D['activated']=True;D['letter']=letter;save()
if 'finance' not in D:
    call({'action':'invoice_save','project_id':pid,'request_id':str(uuid.uuid4())},A,403)
    invoice=call({'action':'invoice_save','project_id':pid,'request_id':str(uuid.uuid4()),'description':'QA fee','invoice_date':'2026-09-01','due_date':'2026-10-01','lines':[{'description':'QA test fee','quantity':1,'unit_price':1000,'vat_rate':16}]},P)['invoice_id'];D['invoice']=invoice;save()
    i=get('invoices','invoice_id',invoice);posted=call({'action':'invoice_post','invoice_id':invoice,'revision':i['revision']},P);assert posted['invoice_number'].startswith('INV/')
    payment=call({'action':'payment_add','invoice_id':invoice,'request_id':str(uuid.uuid4()),'payment_date':'2026-09-05','method':'M-Pesa','reference':'QA'+suffix,'amount':400},P)
    assert payment['receipt_number'].startswith(posted['invoice_number'])
    call({'action':'payment_add','invoice_id':invoice,'request_id':str(uuid.uuid4()),'payment_date':today,'method':'Cash','reference':'QA overpayment','amount':800},P,400)
    i=get('invoices','invoice_id',invoice);call({'action':'invoice_void','invoice_id':invoice,'revision':i['revision'],'reason':'Cannot void allocated payment'},P,400)
    report=call({'action':'statement','client_id':cid,'starts_on':'2026-10-01','ends_on':today},P);assert report['opening']==760 and report['closing']==760
    call({'action':'payment_reverse','payment_id':payment['payment_id'],'request_id':str(uuid.uuid4()),'payment_date':today,'reason':'QA correction'},P)
    report=call({'action':'statement','client_id':cid,'starts_on':'2026-10-01','ends_on':today},P);assert report['opening']==760 and report['closing']==1160
    call({'action':'budget_add','project_id':pid,'agreed_fee':1000,'budgeted_hours':8,'budgeted_cost':500,'effective_from':'2026-09-01','reason':'Fictional QA baseline'},P)
    call({'action':'wip_report'},P)
    aw=w(A)['suite'];assert not aw['invoices'] and not aw['payments'] and not aw['staff_rates'] and not aw['budgets'] and not aw['compliance']
    D['finance']=True;save()
if 'tax' not in D:
    if not any(r['rate_key']=='QA'+suffix.upper() for r in w()['suite']['rates']):call({'action':'rate_save','rate_key':'QA'+suffix,'tax_type':'Corporation tax','label':'FICTIONAL QA ONLY','effective_from':'2026-01-01','effective_to':'2026-12-31','parameters':{'rate':30},'source_url':'https://www.kra.go.ke/','source_note':'Synthetic test configuration, not a statutory rate recommendation'},P)
    rate=next(r for r in w()['suite']['rates'] if r['rate_key']=='QA'+suffix.upper());call({'action':'rate_approve','rate_id':rate['rate_id'],'revision':rate['revision']},A,403)
    if rate['status']=='Draft':call({'action':'rate_approve','rate_id':rate['rate_id'],'revision':rate['revision']},P)
    case=D.get('case') or call({'action':'tax_create','project_id':pid,'kind':'Company','starts_on':'2026-09-01','ends_on':'2026-09-30'},A)['case_id'];D['case']=case;save()
    r=call({'action':'tax_read','case_id':case},A)
    source_task=next(t['task_id'] for t in w()['tasks'] if t['project_id']==pid and t['assigned_to']==ctx['ASSISTANT']['staff_id'])
    source=r['case'].get('source_document_id') or multipart({'client_id':cid,'firm_id':firm,'project_id':pid,'task_id':source_task,'category':'Working Papers'},book([['Account','Value'],['Fictional profit',100000]]),'QA_computation.xlsx',xlsx,P)['document_id']
    if r['case']['status']=='Draft':call({'action':'tax_save','case_id':case,'revision':r['case']['revision'],'rate_id':rate['rate_id'],'source_document_id':source,'inputs':{'accounting_profit':100000,'addbacks':10000,'deductions':5000,'losses':0,'credits':0,'reliefs':0,'return_tied':True}},A)
    r=call({'action':'tax_read','case_id':case},A);assert r['result']['payable']==31500,r['result']
    if r['case']['status']=='Draft':call({'action':'tax_submit','case_id':case,'revision':r['case']['revision']},A)
    r=call({'action':'tax_read','case_id':case},P)
    if r['case']['status']=='Submitted':call({'action':'tax_review','case_id':case,'revision':r['case']['revision'],'status':'Locked','review_note':'QA independent review'},A,403);call({'action':'tax_review','case_id':case,'revision':r['case']['revision'],'status':'Locked','review_note':'QA independent review'},P)
    r=call({'action':'tax_read','case_id':case},P);call({'action':'tax_save','case_id':case,'revision':r['case']['revision'],'inputs':{}},A,400)
    if r['case']['status']=='Locked':call({'action':'tax_file','case_id':case,'revision':r['case']['revision'],'filing_date':today},P)
    export=call({'action':'suite_export','kind':'Tax','record_id':case},P);z=zipfile.ZipFile(io.BytesIO(base64.b64decode(export['base64'])));assert '02_Working_Papers.xlsx' in z.namelist() and any(n.startswith('Attachments/') for n in z.namelist())
    export=call({'action':'suite_export','kind':'KYC','record_id':cid},P);z=zipfile.ZipFile(io.BytesIO(base64.b64decode(export['base64'])));assert {'KYC_Form.pdf','KYC_Form.docx','KYC_Summary.xlsx'}.issubset(z.namelist())
    D['tax']=True;save()
if 'comms_hr' not in D:
    call({'action':'communication_save','client_id':cid,'firm_id':firm,'project_id':pid,'signatory_id':ctx['staff_id'],'request_id':str(uuid.uuid4()),'channel':'Meeting','subject':'QA test meeting','purpose':'Testing','details':'Fictional meeting record','communication_date':today,'meeting_link':'https://meet.google.com/abc-defg-hij'},P)
    call({'action':'clock','operation':'in','location':'Remote','reason':'Fictional QA check'},A)
    call({'action':'clock','operation':'out','break_minutes':0},A)
    leave=call({'action':'leave_request','leave_type':'Study','starts_on':'2027-01-18','ends_on':'2027-01-19','reason':'Fictional QA request'},A)['leave_id'];l=get('leave','leave_id',leave);call({'action':'leave_review','leave_id':leave,'revision':l['revision'],'status':'Approved','review_note':'QA busy-season cover arranged','capacity_acknowledged':True},P)
    credential=call({'action':'credential_add','credential_type':'CPD','description':'Fictional QA training','status':'Passed','issue_date':today,'cpd_hours':2,'evidence_reference':'QA only'},A)['credential_id'];cr=get('credentials','credential_id',credential);call({'action':'credential_verify','credential_id':credential,'revision':cr['revision']},P)
    ticket=call({'action':'ticket_add','project_id':pid,'assigned_to':ctx['SUPERVISOR']['staff_id'],'subject':'QA support','description':'Synthetic regression','category':'Attendance','priority':'High'},A)['ticket_id'];t=call({'action':'ticket_read','ticket_id':ticket},S)['ticket'];call({'action':'ticket_reply','ticket_id':ticket,'revision':t['revision'],'status':'Resolved','message':'QA resolved'},S)
    D['comms_hr']=True;save()
if 'import' not in D:
    name='S7 imported QA '+suffix;data=book([['Name'],[name],[name],['']]);fields={'action':'master_preview','dataset':'Clients','firm_id':firm}
    multipart(fields,data,'QA_client_import.xlsx',xlsx,A,'master',403)
    multipart(fields,data,'QA_client_import.xlsx',xlsx,P,'master')
    fields.update(action='master_commit',column_map=json.dumps({'company_name':1}))
    r=multipart(fields,data,'QA_client_import.xlsx',xlsx,P,'master');assert r['new_count']==1 and r['duplicate_count']==1,r
    r=multipart(fields,data,'QA_client_import.xlsx',xlsx,P,'master');assert r['new_count']==0 and r['duplicate_count']==2,r
    ws=w();imported=next(c for c in ws['clients'] if c['company_name']==name);assert len(imported['missing_fields'])==4
    login='s7.import.'+suffix;staffdata=book([['ID','Login','Name'],[login,login,'S7 imported QA staff '+suffix],[login,login,'S7 imported QA staff '+suffix]])
    sf={'action':'master_commit','dataset':'Staff','firm_id':firm,'column_map':json.dumps({'staff_identifier':1,'login_id':2,'full_name':3})}
    r=multipart(sf,staffdata,'QA_staff_import.xlsx',xlsx,P,'master');assert r['new_count']==1 and r['duplicate_count']==1,r
    ws=w();st=next(x for x in ws['staff'] if x['login_id']==login);assert st['access_status']=='Pending setup'
    D['imported_staff']=st['staff_id'];D['imported_client']=imported['client_id'];D['import']=True;save()
print('PASS complete staging suite: activation, financial integrity, review separation, reports, HR, communications and imports',flush=True)
